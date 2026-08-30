# -*- coding: utf-8 -*-
"""合意传媒 · 从零自建数据全流程测试用例 Excel 生成器"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '合意传媒测试用例-从零流程.xlsx')

HEADERS = ['编号', '优先级', '模块', '场景', '操作步骤', '预期结果', '测试结果', '备注/记录栏']

# ── 全流程：不依赖林女士/陈老板等种子客户，全部自己新建 ──
FLOW_CASES = [
    ('阶段0-环境与数据准备', [
        ('F-00', 'P0', '准备', '环境检查', '确认后端可访问；Pad已装最新sales App；后台可登录', '三端均可打开，无报错', '', ''),
        ('F-01', 'P0', '准备', '约定测试数据', '在「测试数据记录」Sheet填写：客户名、手机号(唯一)、行业/业态', '填写完成，手机号不与现有客户重复', '', '例：测试张老板·蜀味火锅 / 13900001234'),
        ('F-02', 'P0', '后台', '新建客户档案', 'admin登录→客户管理→新增客户：名称/联系人/手机号=上表；行业=餐饮；业态=火锅；经理=阿杰；状态=跟进中', '保存成功；列表可搜到新客户', '', 'Pad端「新建客户」为演示，须后台建'),
        ('F-03', 'P1', '后台', '新建小程序账号(可选)', '用户管理→新增：账号=客户手机号；密码123456；手机号=同客户；审核=是', '用户可登录小程序', '', 'sms开发模式下验证码多为123456'),
        ('F-04', 'P1', '后台', '上传自有素材(可选)', '素材内容→批量上传mp4→填行业餐饮/火锅→上架', 'Pad素材库可见新素材；选片时可喜欢', '', '不执行则选片用系统已有火锅样片'),
    ]),
    ('阶段1-Pad接待与选片', [
        ('F-10', 'P0', 'Pad', '销售登录', 'Pad打开sales→ajie/123456登录', '进入工作台，问候语含「阿杰」', '', ''),
        ('F-11', 'P0', 'Pad', '找到新客户', '客户页→搜索栏输入新客户名或手机号→选中', '右侧显示画像；状态「跟进中」；已选/已交付为0或很少', '', ''),
        ('F-12', 'P0', 'Pad', '开始选片', '点击「开始选片」', '进入选片页，顶栏显示新客户全名', '', ''),
        ('F-13', 'P0', 'Pad', '挑选素材', '点「喜欢」3~5条（建议含厨过程+硬广各至少1条）', '已选计数递增；右侧五维配方柱变化', '', '记录实际喜欢数：____条'),
        ('F-14', 'P1', 'Pad', '筛选素材', '点chips「厨过程」「教知识」等', '列表按类型过滤，可继续喜欢', '', ''),
        ('F-15', 'P0', 'Pad', '生成内容方案', '点「✦ 结束选片并生成方案」', 'Toast「方案已生成，客户进入待付款」；约0.6s后跳回客户页', '', ''),
        ('F-16', 'P0', 'Pad', '核实待付款', '客户页查看该客户', '状态变「待付款」；下一步「提醒客户完成付款」；副文案含「内容方案已确认，等待付款」', '', ''),
    ]),
    ('阶段2-Pad收款与订单', [
        ('F-20', 'P0', 'Pad', '消息待办', '消息页→筛选「成交」或全部', '出现该客户「待付款」卡片，标题含等待付款', '', ''),
        ('F-21', 'P0', 'Pad', '确认收款', '客户页或消息页→「确认收款并生成订单」→弹窗确认', 'Toast含新订单号(格式YJ-MMDD-xxx)；客户变「已成交」', '', '记录订单号：________'),
        ('F-22', 'P0', 'Pad', '订单列表', '订单页→「待拍摄」Tab→搜索新客户名', '出现刚生成的订单；条数=刚才喜欢数(或默认10)', '', ''),
        ('F-23', 'P0', 'Pad', '订单详情', '点击该订单', '显示订单号/套餐/负责人；内容清单条数与选片一致；五维雷达有数据', '', ''),
        ('F-24', 'P1', 'Pad', '记录跟进', '客户页→「记录本次跟进」→填表→保存', '跟进页无语音按钮；保存成功；时间线新增记录', '', ''),
    ]),
    ('阶段3-后台履约', [
        ('F-30', 'P0', '后台', '找到新订单', '订单管理→搜客户名或订单号', '与Pad F-21订单号一致', '', ''),
        ('F-31', 'P0', '后台', '更新制作进度', '点订单→右侧履约：已完成条数=视频总数→状态选「待交付」→保存', '提示保存成功；Pad订单页进度同步', '', ''),
        ('F-32', 'P1', '后台', '标记完成(可选)', '履约区点「→标记已完成」或手动改状态已完成', '订单进「已完成」Tab；客户已交付数增加', '', ''),
        ('F-33', 'P1', '后台', '录入成品(可选)', '当前后台「成品与优质作品」无新增按钮；需开发协助POST /hyDeliverable/save 关联本订单', '小程序「内容」页出现可点成品', '', '无录入则内容页成品区为空属正常'),
    ]),
    ('阶段4-微信小程序', [
        ('F-40', 'P0', '小程序', '登录', '清除小程序缓存→手机号登录(同客户手机)→验证码(开发模式多为123456)', '登录成功；若F-03已建用户则直接进入', '', '须先F-02建客户档案'),
        ('F-41', 'P0', '小程序', '绑定客户', '登录后自动bindByPhone(或进服务/内容页触发)', '绑定到F-02新建客户；无「未绑定服务账号」提示', '', ''),
        ('F-42', 'P0', '小程序', '我的服务', '底部「服务」', '显示当前订单；进度含「方案已确认」等步骤；已完成数与后台一致', '', ''),
        ('F-43', 'P0', '小程序', '我的内容', '底部「内容」', '对标参考区显示方案条数(=选片数)；成品区：有F-33则显示，否则提示暂无', '', ''),
        ('F-44', 'P1', '小程序', '对标参考', '点「查看」对标参考', 'Toast「对标参考为只读，不可下载」', '', ''),
        ('F-45', 'P1', '小程序', '下载标记', '若有成品→点下载图标', 'Toast「已标记下载」；卡片显示已下载', '', '依赖F-33'),
    ]),
    ('阶段5-闭环核对', [
        ('F-50', 'P0', '核对', '客户档案一致', '后台客户管理 vs Pad客户页 vs 小程序服务信息', '名称/手机/行业/成交次数/剩余条数 三端一致', '', ''),
        ('F-51', 'P0', '核对', '订单数据一致', '后台订单 vs Pad订单 vs 小程序服务编号', '订单号、视频总数、完成数一致', '', ''),
        ('F-52', 'P0', '核对', '方案与选片一致', '后台查hy_content_plan(或Pad消息查看方案) vs 选片喜欢列表', '条数与类型分布与选片时一致', '', ''),
        ('F-53', 'P1', '核对', '消息闭环', 'Pad消息页待付款卡片', 'F-21收款后该卡片消失或变已完成', '', ''),
    ]),
]

# ── 可选：不依赖特定客户的功能回归 ──
OPTIONAL_CASES = [
    ('附-通用功能(可选)', [
        ('O-01', 'P1', 'Pad', '素材库入口', '导航点素材库', '进入方形library页；点火锅可进选片', '', ''),
        ('O-02', 'P1', 'Pad', '视频缓存', 'App登录后点工作台「本地x/y」同步', 'Toast素材已同步；选片本地优先播放', '', '仅App-PLUS'),
        ('O-03', 'P1', 'Pad', '竖屏布局', '旋转Pad竖屏浏览各页', '底栏导航可用，布局正常', '', ''),
        ('O-04', 'P2', 'Pad', '无客户结束选片', '素材库直接进选片→不选客户点结束', '提示「请先从客户页进入选片，再生成方案」', '', ''),
    ]),
]

DATA_SHEET_ROWS = [
    ['本次测试数据记录（执行前填写，执行中回填）', '', '', '', ''],
    ['', '', '', '', ''],
    ['字段', '填写说明', '你的填写', '回填(执行后)', ''],
    ['测试批次', '如 2026-08-30-第1轮', '', '', ''],
    ['新客户全称', '例：测试张老板·蜀味火锅', '', '', '须含店名便于搜索'],
    ['联系人', '例：张老板', '', '', ''],
    ['客户手机号', '11位，全局唯一', '', '', '须与小程序登录号一致'],
    ['行业/业态', '建议：餐饮 / 火锅', '', '', ''],
    ['负责经理', '阿杰', '', '', ''],
    ['', '', '', '', ''],
    ['执行记录（流程中回填）', '', '', '', ''],
    ['客户ID', '后台保存后可查', '', '', ''],
    ['选片喜欢数', 'F-13实际选择', '', '', ''],
    ['选片素材ID', '可抄选片页或后台', '', '', '逗号分隔 如7001,7002'],
    ['内容方案ID', 'F-15后后台可查', '', '', ''],
    ['订单号', 'F-21 Toast', '', '', 'YJ-MMDD-xxx'],
    ['订单ID', '后台订单管理', '', '', ''],
    ['收款后客户状态', '应为已成交', '', '', ''],
    ['小程序绑定结果', '服务页客户名', '', '', ''],
    ['', '', '', '', ''],
    ['注意事项', '', '', '', ''],
    ['1', '全流程只用本表客户，不要用林女士/陈老板等种子客户', '', '', ''],
    ['2', '收款(F-21)必须在生成方案(F-15)之后，同一客户连续操作', '', '', ''],
    ['3', '小程序须先有hy_customer(手机一致)，再登录yonghu账号', '', '', ''],
    ['4', '成品播放需F-33录入；无录入时验证对标参考与订单进度即可', '', '', ''],
]

INFO_ROWS = [
    ['合意传媒 · 从零自建数据全流程测试', '', '', '', '', '', '', ''],
    ['', '', '', '', '', '', '', ''],
    ['适用说明', '本用例不依赖种子假客户，请按F系列顺序执行，数据在「测试数据记录」Sheet自建', '', '', '', '', ''],
    ['', '', '', '', '', '', '', ''],
    ['环境', '内容', '', '', '', '', '', ''],
    ['后端', 'https://hycm.baibaiyeye.com.cn/ssm48yhg/', '', '', '', '', '', ''],
    ['Pad账号', 'ajie / 123456', '', '', '', '', '', ''],
    ['后台账号', 'admin / 123456', '', '', '', '', '', ''],
    ['小程序', '手机号=客户手机号；开发环境验证码多为123456', '', '', '', '', '', ''],
    ['', '', '', '', '', '', '', ''],
    ['推荐顺序', '填表→F-00~F-04准备→F-10~F-16选片→F-20~F-24收款→F-30~F-33履约→F-40~F-45小程序→F-50~F-53核对', '', '', '', '', '', ''],
    ['Sheet说明', '全流程主链路=必测；测试数据记录=填空；附-通用功能=有时间再做', '', '', '', '', '', ''],
]


def style_header(ws, row, cols):
    fill = PatternFill('solid', fgColor='2F6BFF')
    font = Font(bold=True, color='FFFFFF', size=11)
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def auto_width(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_cases_sheet(ws, cases, thin_border):
    ws.append(HEADERS)
    style_header(ws, 1, len(HEADERS))
    row_idx = 2
    for section, items in cases:
        ws.cell(row=row_idx, column=1, value=section)
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=8)
        ws.cell(row=row_idx, column=1).font = Font(bold=True, size=11)
        ws.cell(row=row_idx, column=1).fill = PatternFill('solid', fgColor='EAF1FF')
        row_idx += 1
        for case in items:
            ws.append(list(case))
            row_idx += 1
    for r in range(2, row_idx):
        for c in range(1, 9):
            ws.cell(row=r, column=c).alignment = Alignment(vertical='top', wrap_text=True)
            ws.cell(row=r, column=c).border = thin_border
    auto_width(ws, [10, 8, 12, 16, 44, 44, 10, 22])


def main():
    wb = Workbook()
    thin = Side(style='thin', color='E9EDF3')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Sheet1: 说明
    ws0 = wb.active
    ws0.title = '测试说明'
    for r, row in enumerate(INFO_ROWS, 1):
        for c, val in enumerate(row, 1):
            cell = ws0.cell(row=r, column=c, value=val)
            if r == 1:
                cell.font = Font(bold=True, size=14)
    auto_width(ws0, [16, 52, 12, 12, 12, 12, 12, 12])

    # Sheet2: 测试数据记录
    ws_data = wb.create_sheet('测试数据记录', 1)
    for r, row in enumerate(DATA_SHEET_ROWS, 1):
        for c, val in enumerate(row, 1):
            cell = ws_data.cell(row=r, column=c, value=val)
            if r in (1, 11):
                cell.font = Font(bold=True, size=12)
            if r == 1:
                cell.fill = PatternFill('solid', fgColor='FFF3E0')
    auto_width(ws_data, [18, 28, 28, 28, 8])
    # 高亮填写列
    for r in range(4, 10):
        ws_data.cell(row=r, column=3).fill = PatternFill('solid', fgColor='FFFDE7')
    for r in range(12, 20):
        ws_data.cell(row=r, column=4).fill = PatternFill('solid', fgColor='E8F5E9')

    # Sheet3: 全流程主链路
    ws_flow = wb.create_sheet('全流程主链路', 2)
    write_cases_sheet(ws_flow, FLOW_CASES, border)

    # Sheet4: 全部汇总
    ws_all = wb.create_sheet('全部用例', 3)
    write_cases_sheet(ws_all, FLOW_CASES + OPTIONAL_CASES, border)

    # 分阶段 sheet
    for section, items in FLOW_CASES:
        name = section.split('-', 1)[-1][:20]
        for ch in ['/', '\\', '*', '?', ':', '[', ']']:
            name = name.replace(ch, '-')
        ws = wb.create_sheet(name)
        ws.append(HEADERS)
        style_header(ws, 1, len(HEADERS))
        for case in items:
            ws.append(list(case))
        for r in range(2, 2 + len(items)):
            for c in range(1, 9):
                ws.cell(row=r, column=c).alignment = Alignment(vertical='top', wrap_text=True)
                ws.cell(row=r, column=c).border = border
        auto_width(ws, [10, 8, 12, 16, 44, 44, 10, 22])

    # 可选
    ws_opt = wb.create_sheet('附-通用功能')
    ws_opt.append(HEADERS)
    style_header(ws_opt, 1, len(HEADERS))
    for case in OPTIONAL_CASES[0][1]:
        ws_opt.append(list(case))
    for r in range(2, 2 + len(OPTIONAL_CASES[0][1])):
        for c in range(1, 9):
            ws_opt.cell(row=r, column=c).alignment = Alignment(vertical='top', wrap_text=True)
            ws_opt.cell(row=r, column=c).border = border
    auto_width(ws_opt, [10, 8, 12, 16, 44, 44, 10, 22])

    wb.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
